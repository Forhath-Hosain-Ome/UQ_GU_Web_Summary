"""Exercise the existing worker using only generated images and a tiny template."""
import tempfile
from pathlib import Path
from unittest.mock import patch

from django.contrib.auth.models import User
from django.test import TestCase, override_settings
from docx import Document
from PIL import Image

from image_processor.models import FolderBatch, FolderReport, FolderFailedPDF
from image_processor.tasks import process_folder_task


class MultiFolderWorkerTests(TestCase):
    def test_all_folders_and_partial_failure_generate_independent_reports(self):
        for fail_last in (False, True):
            with self.subTest(fail_last=fail_last), tempfile.TemporaryDirectory() as tmp:
                base = Path(tmp)
                root = base / 'job'
                template = base / 'template.docx'
                doc = Document()
                doc.add_table(rows=1, cols=1)
                doc.save(template)
                for name in ('A', 'B', 'C'):
                    folder = root / name / 'nested'
                    folder.mkdir(parents=True)
                    image = folder / 'same.png'
                    if fail_last and name == 'C':
                        image.write_bytes(b'invalid image fixture')
                    else:
                        Image.new('RGB', (8, 8), color='red').save(image)
                batch = FolderBatch.objects.create(
                    created_by=User.objects.create(username=f'user-{fail_last}'),
                    source_folder=str(root), total_folders=3,
                )
                with override_settings(MEDIA_ROOT=base, IMAGE_PROCESSOR_SETTINGS={'OUTPUT_DIR': base / 'output'}):
                    with patch('image_processor.tasks.process_image._push') as push:
                        result = process_folder_task.run(batch.pk, str(root), template_path=str(template))
                batch.refresh_from_db()
                self.assertEqual(result['status'], 'PARTIAL' if fail_last else 'COMPLETED')
                self.assertEqual(batch.total_folders, 3)
                self.assertEqual(batch.processed_folders, 2 if fail_last else 3)
                self.assertEqual(batch.failed_folders, int(fail_last))
                reports = FolderReport.objects.filter(batch=batch)
                self.assertEqual(set(reports.values_list('folder_name', flat=True)), {'A', 'B', 'C'})
                for report in reports.filter(status='COMPLETED'):
                    output = base / report.pdf_output_path
                    self.assertTrue(output.is_file())
                    self.assertEqual(len(Document(output).inline_shapes), 1)
                    self.assertEqual(report.image_count, 1)
                self.assertEqual(FolderFailedPDF.objects.filter(batch=batch).count(), int(fail_last))
                messages = [call.args[1] for call in push.call_args_list]
                self.assertEqual(messages[0]['status'], 'PROCESSING')
                self.assertEqual(messages[-1]['type'], 'batch.complete')
                self.assertEqual(messages[-1]['progress_percent'], 100)
                self.assertEqual(messages[-1]['failed'], int(fail_last))
                self.assertFalse(root.exists())
