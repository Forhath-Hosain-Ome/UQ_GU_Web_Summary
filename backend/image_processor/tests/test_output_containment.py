"""Generated fixtures only: output boundaries, API parity and worker persistence."""
import io
import tempfile
import zipfile
from pathlib import Path
from unittest.mock import patch

from django.contrib.auth.models import User
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import SimpleTestCase, TestCase, override_settings
from docx import Document
from PIL import Image

from image_processor.models import FolderBatch, FolderFailedPDF
from image_processor.output_paths import (
    OutputPathError, open_output_file, output_directory, output_path,
    report_filename, validate_report_date,
)
from image_processor.tasks import process_folder_task
from image_processor.tests import test_upload_containment as upload_tests

LABEL = "Defect_GMTS_pictures_report_for_Style"
BAD_DATES = (
    "../../../../escaped", r"..\..\outside", "/absolute", r"C:\outside",
    "2026-02-29", "0000-01-01", "10000-01-01", "2026-13-01", "2026-01-00",
    "20260908", "2026-9-8", "2026-09-08T00:00:00", "2026-09-08/evil",
    "2026-09-08\x00", "2026-09\n-08", "２０２６-09-08", "x" * 1025,
)


class OutputRulesTests(SimpleTestCase):
    def test_calendar_date_boundaries_and_invalid_types(self):
        for value in ("", "0001-01-01", "9999-12-31", "2024-02-29", "2000-02-29"):
            self.assertEqual(validate_report_date(value), value)
        self.assertEqual(validate_report_date(" 2026-09-08 "), "2026-09-08")
        self.assertEqual(validate_report_date("  "), "")
        for value in (*BAD_DATES, None, 20260908, [], {}, True, "1900-02-29"):
            with self.subTest(value=value), self.assertRaises(OutputPathError):
                validate_report_date(value)

    def test_filename_contract_and_utf8_byte_boundary(self):
        self.assertEqual(report_filename("Style A", "", LABEL), f"{LABEL}_Style_A_dated on no-date.docx")
        overhead = len(report_filename("A", "2026-09-08", "L").encode()) - 1
        for folder in ("a" * (255 - overhead), "é" * ((255 - overhead) // 2)):
            self.assertLessEqual(len(report_filename(folder, "2026-09-08", "L").encode()), 255)
        with self.assertRaises(OutputPathError):
            report_filename("a" * (256 - overhead), "2026-09-08", "L")
        for component in ("..", "a/b", r"a\b", "C:foo", "CON", "end.", "end ", "a\x00b", "", None):
            for folder, label in ((component, "L"), ("A", component)):
                with self.subTest(folder=folder, label=label), self.assertRaises(OutputPathError):
                    report_filename(folder, "", label)

    def test_containment_exclusive_write_and_generic_errors(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp).resolve()
            root = base / "output"
            output_directory(root, "1/A")
            target = root / "1/A/report.docx"
            with open_output_file(root, "1/A/report.docx") as stream:
                stream.write(b"original")
            with self.assertRaisesRegex(OutputPathError, "^Invalid report output path.$"):
                with open_output_file(root, "1/A/report.docx"):
                    self.fail("Existing target opened")
            self.assertEqual(target.read_bytes(), b"original")
            for path in ("../escape", "/absolute", r"C:\outside", "1/../../outside", "1/A/../escape"):
                with self.subTest(path=path), self.assertRaises(OutputPathError):
                    output_path(root, path)
            real_resolve = Path.resolve
            with patch.object(Path, "resolve", lambda path, *a, **kw: base / "outside" if path.name == "new.docx" else real_resolve(path, *a, **kw)):
                with self.assertRaises(OutputPathError):
                    output_path(root, "1/A/new.docx")
            self.assertFalse((base / "outside").exists())

    def test_real_directory_and_file_links_cannot_redirect_output(self):
        with tempfile.TemporaryDirectory() as tmp:
            base = Path(tmp).resolve()
            root = base / "output"
            root.mkdir()
            outside = base / "outside"
            outside.mkdir()
            sentinel = outside / "sentinel.docx"
            sentinel.write_bytes(b"keep")
            try:
                (root / "linked").symlink_to(outside, target_is_directory=True)
                (root / "file.docx").symlink_to(sentinel)
            except OSError as exc:
                self.skipTest(f"Symlink creation unavailable: {exc}")
            for relative in ("linked/new.docx", "file.docx"):
                with self.subTest(relative=relative), self.assertRaises(OutputPathError):
                    with open_output_file(root, relative):
                        self.fail("Link followed")
            self.assertEqual(sentinel.read_bytes(), b"keep")
            self.assertEqual(list(outside.iterdir()), [sentinel])


class OutputDateApiTests(TestCase):
    setUp = upload_tests.UploadContainmentTests.setUp

    def submit(self, date, archive):
        data = {"date": date}
        if archive:
            buffer = io.BytesIO()
            with zipfile.ZipFile(buffer, "w") as zf:
                zf.writestr("A/photo.png", b"fixture")
            data["archive"] = SimpleUploadedFile("images.zip", buffer.getvalue())
        else:
            data["files"] = [SimpleUploadedFile("photo.png", b"fixture")]
            data["paths"] = ["A/photo.png"]
        with self.captureOnCommitCallbacks(execute=True):
            return self.client.post("/image/folder/upload/", data, format="multipart")

    def test_invalid_date_rejected_before_staging_in_both_modes(self):
        for archive in (False, True):
            for date in BAD_DATES:
                with self.subTest(archive=archive, date=date):
                    response = self.submit(date, archive)
                    self.assertEqual(response.status_code, 400)
                    self.assertEqual(response.data["code"], "invalid_date")
                    self.assertEqual(FolderBatch.objects.count(), 0)
                    self.assertFalse(self.staging.exists())
                    self.delay.assert_not_called()

    def test_valid_and_blank_dates_preserve_dispatch_contract_in_both_modes(self):
        for archive in (False, True):
            for date in ("", "  ", " 2024-02-29 "):
                with self.subTest(archive=archive, date=date):
                    response = self.submit(date, archive)
                    self.assertEqual(response.status_code, 201, response.data)
                    self.assertEqual(self.delay.call_args.args[2], date.strip())


class OutputWorkerTests(TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.addCleanup(self.tmp.cleanup)
        self.base = Path(self.tmp.name).resolve()
        self.root = self.base / "job"
        self.output = self.base / "output"
        self.template = self.base / "template.docx"
        doc = Document()
        doc.add_table(rows=1, cols=1)
        doc.save(self.template)
        self.user = User.objects.create(username="output-owner")

    def batch(self, names=("A",)):
        for name in names:
            folder = self.root / name
            folder.mkdir(parents=True, exist_ok=True)
            Image.new("RGB", (8, 8), color="blue").save(folder / "photo.png")
        return FolderBatch.objects.create(created_by=self.user, source_folder=str(self.root), total_folders=len(names))

    def run_worker(self, batch, **kwargs):
        with override_settings(MEDIA_ROOT=self.base, IMAGE_PROCESSOR_SETTINGS={"OUTPUT_DIR": self.output}):
            with patch("image_processor.tasks.process_image._push"):
                return process_folder_task.run(batch.pk, str(self.root), template_path=str(self.template), **kwargs)

    def test_baseline_traversal_and_invalid_worker_inputs_fail_without_output(self):
        for date in (*BAD_DATES, None, 123, []):
            with self.subTest(date=date):
                batch = self.batch()
                result = self.run_worker(batch, date=date)
                batch.refresh_from_db()
                self.assertEqual(result["status"], "FAILED")
                self.assertEqual(batch.failed_folders, 1)
                report = batch.reports.get()
                self.assertEqual(report.pdf_output_path, "")
                self.assertEqual(report.error_message, "Invalid report date.")
                self.assertEqual(FolderFailedPDF.objects.filter(batch=batch).count(), 1)
                self.assertFalse(self.output.exists())
                self.assertFalse(self.root.exists())

    def test_real_docx_paths_contents_and_blank_date_compatibility(self):
        for date in ("", "2024-02-29"):
            with self.subTest(date=date):
                batch = self.batch(("Style A", "B"))
                self.assertEqual(self.run_worker(batch, date=date)["status"], "COMPLETED")
                for report in batch.reports.all():
                    expected = self.output / str(batch.pk) / report.folder_name / report_filename(report.folder_name, date, LABEL)
                    self.assertEqual((self.base / report.pdf_output_path).resolve(), expected)
                    self.assertEqual(len(Document(expected).inline_shapes), 1)
                    self.assertFalse(Path(report.pdf_output_path).is_absolute())

    def test_invalid_label_fails_before_output_creation(self):
        batch = self.batch()
        self.assertEqual(self.run_worker(batch, label_type="../../escape")["status"], "FAILED")
        self.assertFalse(self.output.exists())

    def test_existing_docx_and_prepared_files_survive_failed_writes(self):
        for prepared in (False, True):
            with self.subTest(prepared=prepared):
                batch = self.batch()
                folder = self.output / str(batch.pk) / "A"
                target = folder / ("_prepared/photo.jpg" if prepared else report_filename("A", "", LABEL))
                target.parent.mkdir(parents=True)
                target.write_bytes(b"existing-report")
                self.assertEqual(self.run_worker(batch)["status"], "FAILED")
                self.assertEqual(target.read_bytes(), b"existing-report")
                self.assertEqual(batch.reports.get().pdf_output_path, "")

    def test_linked_folder_fails_but_other_folder_completes(self):
        batch = self.batch(("A", "B"))
        outside = self.base / "outside"
        outside.mkdir()
        batch_output = self.output / str(batch.pk)
        batch_output.mkdir(parents=True)
        try:
            (batch_output / "A").symlink_to(outside, target_is_directory=True)
        except OSError as exc:
            self.skipTest(f"Symlink creation unavailable: {exc}")
        self.assertEqual(self.run_worker(batch)["status"], "PARTIAL")
        self.assertEqual(list(outside.iterdir()), [])
        self.assertEqual(batch.reports.get(folder_name="A").pdf_output_path, "")
        self.assertTrue((self.base / batch.reports.get(folder_name="B").pdf_output_path).is_file())
