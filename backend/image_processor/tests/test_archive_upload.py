import io
import os
import stat
import struct
import zipfile
from pathlib import Path
from unittest.mock import patch
from unittest import skipUnless

from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase, override_settings
from docx import Document
from PIL import Image
from image_processor.tasks import process_folder_task
from image_processor.models import FolderBatch
from image_processor.tests import test_upload_containment as containment

ARCHIVE = 'image_processor.archive_upload'


def make_zip(entries, compression=zipfile.ZIP_STORED):
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, 'w', compression=compression) as archive:
        for name, data in entries:
            archive.writestr(name, data)
    return buffer.getvalue()


class ArchiveUploadTests(TestCase):
    setUp = containment.UploadContainmentTests.setUp

    def upload_archive(self, entries=None, payload=None, **fields):
        if payload is None:
            payload = make_zip(entries if entries is not None else [('A/photo.jpg', b'image')])
        data = {'archive': SimpleUploadedFile('anything.bin', payload, 'application/octet-stream'), **fields}
        with self.captureOnCommitCallbacks(execute=True):
            return self.client.post('/image/folder/upload/', data, format='multipart')

    def rejected(self, response):
        self.assertEqual(response.status_code, 400, response.data)
        self.assertFalse(FolderBatch.objects.exists())
        self.delay.assert_not_called()
        self.assertEqual(list(self.staging.iterdir()) if self.staging.exists() else [], [])
        self.assertIn('code', response.data)

    def test_success_contract_one_multiple_nested_and_distinct_repeated_names(self):
        for paths in (['A/photo.jpg'], ['A/photo.jpg', 'B/photo.jpg'], ['A/nested/photo.jpg', 'B/photo.jpg']):
            with self.subTest(paths=paths):
                response = self.upload_archive([(p, p.encode()) for p in paths])
                self.assertEqual(response.status_code, 201, response.data)
                batch = FolderBatch.objects.get(pk=response.data['batch_id'])
                root = Path(batch.source_folder)
                self.assertEqual(root.parent, self.staging)
                self.assertEqual(batch.created_by, self.user)
                self.assertEqual(batch.total_folders, len({p.split('/')[0] for p in paths}))
                for path in paths:
                    self.assertEqual((root / path).read_bytes(), path.encode())
                self.delay.assert_called_once_with(batch.pk, str(root), '', is_renamed_file=False)
                self.delay.reset_mock()

    def test_flat_archive_uses_default_or_style_and_normalizes_names(self):
        for style in ('', 'Style-A'):
            response = self.upload_archive([('image one.JPG', b'image')], style=style)
            self.assertEqual(response.status_code, 201, response.data)
            batch = FolderBatch.objects.get(pk=response.data['batch_id'])
            self.assertEqual(batch.total_folders, 1)
            self.assertTrue((Path(batch.source_folder) / (style or 'uploaded') / 'image_one.jpg').is_file())

    def test_directory_records_are_validated_but_empty_folders_do_not_count(self):
        response = self.upload_archive([('A/', b''), ('A/nested/', b''), ('A/nested/a.jpg', b'image'), ('Empty/', b'')])
        self.assertEqual(response.status_code, 201, response.data)
        self.assertEqual(FolderBatch.objects.get().total_folders, 1)

    def test_mixed_separators_are_canonical(self):
        response = self.upload_archive([('A\\nested/a.jpg', b'image')])
        self.assertEqual(response.status_code, 201, response.data)
        self.assertTrue((Path(FolderBatch.objects.get().source_folder) / 'A/nested/a.jpg').is_file())

    def test_ambiguous_modes_are_rejected(self):
        for fields in ({'files': SimpleUploadedFile('a.jpg', b'image')}, {'files': ''}, {'paths': ['A/a.jpg']}, {'paths[]': ['A/a.jpg']}):
            self.rejected(self.upload_archive(**fields))
        data = {'archive': [SimpleUploadedFile('a.zip', b'zip'), SimpleUploadedFile('b.zip', b'zip')]}
        self.rejected(self.client.post('/image/folder/upload/', data, format='multipart'))
        self.rejected(self.client.post('/image/folder/upload/', {'archive': 'not a file'}, format='multipart'))

    def test_empty_corrupt_truncated_and_non_zip(self):
        for payload in (make_zip([]), b'not zip', make_zip([('A/a.jpg', b'image')])[:-8]):
            self.rejected(self.upload_archive(payload=payload))

    def test_no_images_unsupported_empty_image_and_mixed_root_entries(self):
        for entries in ([('A/', b'')], [('A/readme.txt', b'text')], [('A/a.jpg', b'')],
                        [('A/a.jpg', b'image'), ('A/readme.txt', b'text')],
                        [('a.jpg', b'image'), ('A/b.jpg', b'image')]):
            self.rejected(self.upload_archive(entries))

    def test_path_attack_matrix(self):
        for path in ('../a.jpg', '../../a.jpg', 'A/../../../a.jpg', '/a.jpg', 'C:/a.jpg',
                     'C:\\a.jpg', 'C:a.jpg', '\\\\server\\share\\a.jpg', 'A\\..\\a.jpg',
                     'A//a.jpg', './a.jpg', 'NUL/a.jpg', 'A:x/a.jpg', 'A./a.jpg',
                     'A /a.jpg', 'A\x01/a.jpg', 'x' * 256 + '/a.jpg', 'A/' * 520 + 'a.jpg'):
            with self.subTest(path=path):
                response = self.upload_archive([(path, b'image')])
                self.rejected(response)
                self.assertNotIn(path, str(response.data))

    def test_nul_name_is_rejected_before_zipfile_truncation(self):
        payload = make_zip([('A/Xa.jpg', b'image')]).replace(b'A/Xa.jpg', b'A/\x00a.jpg')
        self.rejected(self.upload_archive(payload=payload))

    def test_links_special_files_and_dos_reparse_attributes(self):
        for attr in (stat.S_IFLNK << 16, stat.S_IFCHR << 16, stat.S_IFIFO << 16, 0x400, 0x08, 0x10):
            entry = zipfile.ZipInfo('A/a.jpg')
            entry.create_system = 3
            entry.external_attr = attr
            self.rejected(self.upload_archive([(entry, b'image')]))

    def test_encryption_and_unsupported_compression(self):
        for flags, method in ((1, 0), (0, 99)):
            payload = bytearray(make_zip([('A/a.jpg', b'image')]))
            central = payload.index(b'PK\x01\x02')
            struct.pack_into('<HH', payload, 6, flags, method)
            struct.pack_into('<HH', payload, central + 8, flags, method)
            self.rejected(self.upload_archive(payload=bytes(payload)))

    def test_collisions_preflight_before_writing(self):
        for paths in (['A/a.jpg', 'A/a.jpg'], ['A/a.jpg', 'A/A.jpg'],
                      ['A/a b.jpg', 'A/a_b.jpg'], ['A/a.jpg', 'a/b.jpg'],
                      ['A.jpg', 'A.jpg/b.jpg'], ['A.jpg', 'A.jpg/']):
            with patch(ARCHIVE + '.open_upload_file') as writer:
                self.rejected(self.upload_archive([(p, b'' if p.endswith('/') else b'image') for p in paths]))
                writer.assert_not_called()

    def test_metadata_resource_limits(self):
        for constant, limit, entries in (
            ('MAX_ARCHIVE_SIZE', 10, [('A/a.jpg', b'image')]),
            ('MAX_ENTRIES', 1, [('A/a.jpg', b'image'), ('B/a.jpg', b'image')]),
            ('MAX_FILE_SIZE', 4, [('A/a.jpg', b'image')]),
            ('MAX_TOTAL_SIZE', 8, [('A/a.jpg', b'image'), ('B/a.jpg', b'image')]),
            ('MAX_DIRECTORY_SIZE', 1, [('A/a.jpg', b'image')]),
        ):
            with self.subTest(constant=constant), patch(ARCHIVE + '.' + constant, limit):
                self.rejected(self.upload_archive(entries))
        self.rejected(self.upload_archive(payload=make_zip([('A/a.jpg', b'0' * 10000)], zipfile.ZIP_DEFLATED)))

    def test_small_deflated_archive_is_accepted(self):
        response = self.upload_archive(payload=make_zip([('A/a.jpg', bytes(range(256)))], zipfile.ZIP_DEFLATED))
        self.assertEqual(response.status_code, 201, response.data)

    def test_crc_failure_after_first_file_cleans_without_batch_or_dispatch(self):
        payload = make_zip([('A/first.jpg', b'first'), ('B/second.jpg', b'second')])
        payload = payload.replace(b'secondPK', b'brokenPK', 1)
        self.rejected(self.upload_archive(payload=payload))

    def test_runtime_expansion_limit_is_enforced(self):
        with patch.object(zipfile.ZipExtFile, 'read', return_value=b'x' * 10):
            self.rejected(self.upload_archive())

    def test_zip_library_failure_redacts_exception(self):
        secret = str(self.sandbox / 'private-uuid')
        payload = make_zip([('A/a.jpg', b'image')])
        with patch.object(zipfile.ZipFile, 'open', side_effect=zipfile.BadZipFile(secret)):
            response = self.upload_archive(payload=payload)
        self.rejected(response)
        self.assertNotIn(secret, str(response.data))

    def test_zip_intake_feeds_real_existing_worker(self):
        image = io.BytesIO()
        Image.new('RGB', (8, 8), color='blue').save(image, 'PNG')
        response = self.upload_archive([(p, image.getvalue()) for p in ('A/nested/a.png', 'B/a.png')])
        self.assertEqual(response.status_code, 201, response.data)
        batch = FolderBatch.objects.get()
        template = self.sandbox / 'template.docx'
        doc = Document()
        doc.add_table(rows=1, cols=1)
        doc.save(template)
        with override_settings(MEDIA_ROOT=self.sandbox, IMAGE_PROCESSOR_SETTINGS={'OUTPUT_DIR': self.sandbox / 'output'}):
            with patch('image_processor.tasks.process_image._push'):
                result = process_folder_task.run(batch.pk, batch.source_folder, template_path=str(template))
        self.assertEqual(result['processed'], 2)
        self.assertEqual(result['status'], 'COMPLETED')
        for report in batch.reports.all():
            self.assertEqual(len(Document(self.sandbox / report.pdf_output_path).inline_shapes), 1)

    def test_bad_style_and_unauthenticated_archive_are_rejected(self):
        self.rejected(self.upload_archive(style='../escape'))
        self.client = containment.APIClient()
        response = self.upload_archive()
        self.assertEqual(response.status_code, 401)
        self.assertFalse(FolderBatch.objects.exists())
        self.delay.assert_not_called()

    def test_zip64_multidisk_and_directory_budget_are_rejected_before_parsing(self):
        for field_offset, value in ((4, 1), (10, 65535), (10, 1001)):
            payload = bytearray(make_zip([('A/a.jpg', b'image')]))
            end = payload.rfind(b'PK\x05\x06')
            struct.pack_into('<H', payload, end + field_offset, value)
            with patch(ARCHIVE + '.zipfile.ZipFile') as parser:
                self.rejected(self.upload_archive(payload=bytes(payload)))
                parser.assert_not_called()

    @skipUnless(os.name == 'nt', 'Windows canonical namespace spelling')
    def test_server_resolved_namespace_alias_does_not_false_reject(self):
        real_resolve = Path.resolve
        def extended(path, *args, **kwargs):
            resolved = real_resolve(path, *args, **kwargs)
            return Path('\\\\?\\' + str(resolved)) if path.name == 'a.jpg' else resolved
        with patch.object(Path, 'resolve', extended):
            response = self.upload_archive([('A/a.jpg', b'image')])
        self.assertEqual(response.status_code, 201, response.data)
        def escaped(path, *args, **kwargs):
            return Path('\\\\?\\' + str(self.sandbox / 'outside/a.jpg')) if path.name == 'a.jpg' else real_resolve(path, *args, **kwargs)
        with patch.object(Path, 'resolve', escaped):
            response = self.upload_archive([('A/a.jpg', b'image')])
        self.assertEqual(response.status_code, 400, response.data)
        self.assertFalse((self.sandbox / 'outside').exists())
        self.assertEqual(FolderBatch.objects.count(), 1)
        self.delay.assert_called_once()
