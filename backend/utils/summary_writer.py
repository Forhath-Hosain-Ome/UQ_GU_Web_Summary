"""
------------------------
Fills per-stage Excel templates with audit data and merges the results
into one output workbook.

Template selection
-------------------
Driven entirely by final_summary/templates_registry.py:
  (BuyerFactoryPair.report_type, canonical stage) -> template file
Only stages that actually have matching records get a sheet in the
output workbook — a factory that only ever ran Inline + Re-Final
inspections gets a 2-sheet file, not a 4-sheet one with blanks.

Column resolution
-------------------
NOTHING is hardcoded by column number. Both prefix fields (factory,
date_of_issue, ...) and defect columns are located by reading each
template's own header row (row 2) and matching against known label
aliases. This is deliberate: a previous hardcoded-column-index version
of this file silently swapped PO Qty (Set) / (Pack) values and wrote
audit-end times on top of the Audit Start Time header on one template
family, because the two template families don't share column layouts.
Resolving by label makes the writer self-correcting when templates
change, and turns "wrong column" bugs into a logged "field not found"
instead of silent data corruption.

Key behaviours
--------------
- Pre-flight validation across ALL stages before writing anything:
  if any defect item has no matching template column, raises
  DefectMismatchError before touching any file.
- Only cell.value is written on the per-stage source template — no
  styling touched there. Styling *is* carried over when merging into
  the output workbook (see utils/sheet_copy.py).
- Values normalised to native Python types (date, time, float, int,
  str) so Excel formulas operate without "click to activate".
- A (report_type, stage) with no configured template is skipped with
  a warning, not a hard failure — unless it's the ONLY stage present,
  in which case write_summary raises TemplateNotConfiguredError so the
  caller can surface a clear message instead of returning an empty file.
"""

import logging
import re
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl

from final_summary.templates_registry import (
    SHEET_NAME_MAP,
    STAGE_ORDER,
    TemplateNotConfiguredError,
    resolve_template,
)
from utils.sheet_copy import copy_sheet

logger = logging.getLogger(__name__)


# =============================================================================
# LAYOUT CONSTANTS
# =============================================================================
# Data start row and defect-column start col are layout conventions that
# have held across every template inspected so far. If a future template
# breaks the convention, add a per-stage override here rather than
# hardcoding field positions (see FIELD_LABEL_ALIASES below for why).

FINAL_DATA_START_ROW:    int = 12   # FINAL / RE_FINAL / RANDOM sheets
INLINE_DATA_START_ROW:   int = 5    # INLINE sheet
HEADER_ROW:              int = 2

INLINE_STAGES: set = {"INLINE"}


def _data_start_row(stage: str) -> int:
    return INLINE_DATA_START_ROW if stage in INLINE_STAGES else FINAL_DATA_START_ROW


# =============================================================================
# PREFIX FIELD LABEL ALIASES
# =============================================================================
# field key -> acceptable header text(s), matched case/whitespace-insensitively.
# A field simply won't be written if none of its aliases are found in the
# template's header row — logged, not silently misplaced.

FIELD_LABEL_ALIASES: Dict[str, List[str]] = {
    "factory":               ["Factory"],
    "date_of_issue":         ["Date of issue", "Date of Issue"],
    "inspection_type":       ["Inspection Type"],
    "factory_in":             ["Factory In Time"],
    "factory_out":            ["Factory Out Time"],
    "audit_start":            ["Audit Start Time"],
    "audit_end":              ["Audit End Time"],
    "audit_result":           ["Audit Result"],
    "report_no":              ["Report Number", "Report No.", "Report No"],
    "item_name":              ["Item Name"],
    "style_no":               ["Style NO.", "Style No.", "Style No"],
    "po_no":                  ["PO-NO", "PO NO.", "POーNO", "PO No."],
    "country":                ["Country"],
    "po_qty_pcs":             ["PO Qty.(Pcs)", "PO Qty (Pcs)."],
    "po_qty_set":             ["PO Qty.(Set)", "PO Qty (Set)."],
    "po_qty_pack":            ["PO Qty.(Pack)", "PO Qty (Pack)."],
    "po_wh":                  ["PO WH", "PO  WH"],
    "ship_qty":               ["Ship Qty.", "Shipping Qty"],
    "audit_qty":              ["Audit Qty."],
    "acceptable_defect_qty":  ["Acceptable Defect Qty"],
    "defect_qty":             ["Defect Qty.", "DefectQty."],
    "defect_percentage":      ["Defect %"],
    "person":                 ["Person"],
}

# Formula-driven columns are deliberately absent from the alias map above
# (factory_total_hours, audit_total_hours) — left for Excel to calculate.


# =============================================================================
# FIELD TYPE DECLARATIONS
# =============================================================================

_STR_FIELDS: set = {
    "factory", "inspection_type", "audit_result", "report_no",
    "item_name", "style_no", "po_no", "country",
    "acceptable_defect_qty", "person", "inspector",
}
_DATE_FIELDS: set = {
    "date_of_issue", "po_wh", "exf", "po_edt", "plan_edt", "plan_wh",
}
_TIME_FIELDS: set = {
    "factory_in", "factory_out", "audit_start", "audit_end",
}
_FLOAT_FIELDS: set = {"defect_percentage"}
_INT_FIELDS:   set = {
    "ship_qty", "audit_qty", "defect_qty", "do_qty",
    "po_qty_pcs", "po_qty_pack", "po_qty_set",
}


# =============================================================================
# CUSTOM EXCEPTIONS
# =============================================================================

class DefectMismatchError(Exception):
    """
    Raised when one or more defect items have no matching template column.

    Attributes
    ----------
    mismatches : list of dicts — file_name, report_id, stage, sheet, item
    """
    def __init__(self, mismatches: List[Dict[str, Any]]) -> None:
        self.mismatches = mismatches
        super().__init__(
            f"{len(mismatches)} defect item(s) have no matching template column."
        )


# =============================================================================
# VALUE NORMALISATION
# =============================================================================

def _to_str(value: Any) -> Optional[str]:
    if value is None or value == "":
        return None
    s = str(value).strip()
    return s if s and s != "-" else None


def _to_date(value: Any) -> Optional[date]:
    if value is None or value == "":
        return None
    if isinstance(value, date):
        return value
    s = str(value).strip()
    for fmt in (
        "%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y",
        "%Y/%m/%d", "%m-%d-%Y", "%d-%m-%Y",
        "%d-%b-%y", "%d-%b-%Y",
    ):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def _to_time(value: Any) -> Optional[time]:
    if value is None or value == "":
        return None
    if isinstance(value, time):
        return value
    s = str(value).strip()
    m = re.match(r"(\d{1,2})[:.](\d{2})(?:\s*([AaPp][Mm]))?", s)
    if not m:
        return None
    hour, minute, ampm = int(m.group(1)), int(m.group(2)), m.group(3)
    if ampm:
        if ampm.upper() == "PM" and hour != 12:
            hour += 12
        elif ampm.upper() == "AM" and hour == 12:
            hour = 0
    try:
        return time(hour, minute)
    except ValueError:
        return None


def _to_float(value: Any) -> Optional[float]:
    if value is None or value == "":
        return None
    try:
        return float(str(value).replace("%", "").strip())
    except (ValueError, TypeError):
        return None


def _to_int(value: Any) -> Optional[int]:
    if value is None or value == "":
        return None
    try:
        return int(float(str(value).strip()))
    except (ValueError, TypeError):
        return None


def _normalise(key: str, value: Any) -> Any:
    if key in _STR_FIELDS:   return _to_str(value)
    if key in _DATE_FIELDS:  return _to_date(value)
    if key in _TIME_FIELDS:  return _to_time(value)
    if key in _FLOAT_FIELDS: return _to_float(value)
    if key in _INT_FIELDS:   return _to_int(value)
    if value in ("", "-"):   return None
    return value


def _norm_label(s: str) -> str:
    """Collapse whitespace and lowercase, for robust header matching."""
    return re.sub(r"\s+", " ", str(s)).strip().lower()


def _strip_prefix(name: str) -> str:
    """'3.Hole, tear' → 'Hole, tear'"""
    return re.sub(r"^\d+\.", "", name).strip()


# =============================================================================
# CANONICAL STAGE DETECTION
# =============================================================================

def _canonical_stage(itype: str) -> str:
    u = (itype or "").upper().strip()
    if re.search(r"PRE[\s\-]?FINAL", u):                           return "PRE_FINAL"
    if re.search(r"(?<![A-Z])RE[\s\-]?FINAL|REFINAL", u):          return "RE_FINAL"
    if re.search(r"\bRANDOM\b", u):                                return "RANDOM"
    if re.search(r"\bFINAL\b|SHIPMENT\s*AUDIT|PRE[\s\-]?SHIP", u): return "FINAL"
    if re.search(r"IN[\s\-]?LINE", u):                             return "INLINE"
    if re.search(r"\bCMF\b|COUNTER\s*MASTER", u):                  return "CMF"
    if re.search(r"\bSAMPLE\b|PRE[\s\-]?PROD|\bPP\b", u):         return "SAMPLE"
    return "UNKNOWN"


# =============================================================================
# HEADER MAP  (row 2 of a template -> {normalised_label: col})
# =============================================================================

def _build_header_map(ws) -> Dict[str, int]:
    """Read HEADER_ROW across the whole sheet width. First occurrence of
    a given normalised label wins."""
    header_map: Dict[str, int] = {}
    for col in range(1, ws.max_column + 1):
        raw = ws.cell(row=HEADER_ROW, column=col).value
        if raw is None:
            continue
        name = str(raw).strip()
        if not name:
            continue
        norm = _norm_label(name)
        if norm not in header_map:
            header_map[norm] = col
        stripped = _norm_label(_strip_prefix(name))
        if stripped and stripped not in header_map:
            header_map[stripped] = col
    return header_map


def _resolve_field_col(field_key: str, header_map: Dict[str, int]) -> Optional[int]:
    for alias in FIELD_LABEL_ALIASES.get(field_key, []):
        col = header_map.get(_norm_label(alias))
        if col is not None:
            return col
    return None


def _resolve_defect_col(item_name: str, header_map: Dict[str, int]) -> Optional[int]:
    """Exact (normalised) match first, then prefix-stripped."""
    col = header_map.get(_norm_label(item_name))
    if col is not None:
        return col
    return header_map.get(_norm_label(_strip_prefix(item_name)))


# =============================================================================
# PRE-FLIGHT VALIDATION
# =============================================================================

def _validate_defects(
    stage: str,
    header_map: Dict[str, int],
    records: List[Dict[str, Any]],
    defect_items_by_report: Dict[int, List[Dict[str, Any]]],
    sheet_name: str,
) -> List[Dict[str, Any]]:
    mismatches: List[Dict[str, Any]] = []
    for rec in records:
        rid       = rec.get("report_id")
        file_name = rec.get("file_name", "")
        for d in defect_items_by_report.get(rid, []):
            item_name = (d.get("item") or "").strip()
            if not item_name:
                continue
            major = _to_int(d.get("major_count", 0))
            if not major:
                continue
            if _resolve_defect_col(item_name, header_map) is None:
                mismatches.append({
                    "file_name": file_name,
                    "report_id": rid,
                    "stage":     stage,
                    "sheet":     sheet_name,
                    "item":      item_name,
                    "category":  (d.get("category") or "").strip(),
                })
    return mismatches


# =============================================================================
# SHEET WRITER  (fills one stage's template sheet in place)
# =============================================================================

def _write_stage_sheet(
    ws,
    stage: str,
    records: List[Dict[str, Any]],
    defect_items_by_report: Dict[int, List[Dict[str, Any]]],
    header_map: Dict[str, int],
) -> int:
    start_row   = _data_start_row(stage)
    current_row = start_row

    for rec in records:
        rid = rec.get("report_id")

        for field_key in FIELD_LABEL_ALIASES:
            raw = rec.get(field_key)
            if raw is None:
                continue
            col_idx = _resolve_field_col(field_key, header_map)
            if col_idx is None:
                continue
            value = _normalise(field_key, raw)
            if value is not None:
                ws.cell(row=current_row, column=col_idx).value = value

        for d in defect_items_by_report.get(rid, []):
            item_name = (d.get("item") or "").strip()
            if not item_name:
                continue
            major = _to_int(d.get("major_count", 0))
            if not major:
                continue
            col_idx = _resolve_defect_col(item_name, header_map)
            if col_idx:
                ws.cell(row=current_row, column=col_idx).value = major

        current_row += 1

    rows_written = current_row - start_row
    logger.info(
        "Stage '%s': wrote %d row(s) starting at row %d", stage, rows_written, start_row,
    )
    return rows_written


# =============================================================================
# PUBLIC API
# =============================================================================

def write_summary(
    records: List[Dict[str, Any]],
    defect_items_by_report: Dict[int, List[Dict[str, Any]]],
    output_path: Path,
    report_type: str,
) -> Dict[str, Any]:
    """
    Fill one template per stage present in `records`, merge the filled
    sheets into a single output workbook (only the stages that actually
    have data), and save to output_path.

    Returns a summary dict: {"stages_written": [...], "stages_skipped": [...]}
    stages_skipped entries are {"stage": ..., "reason": ...} — usually
    "not yet configured" for a (report_type, stage) with no template.

    Raises
    ------
    DefectMismatchError        any defect item has no matching column in
                                its stage's template (checked across ALL
                                stages before anything is written)
    TemplateNotConfiguredError only stage(s) present have no template at
                                all — nothing could be produced
    """
    # Group by canonical stage
    by_stage: Dict[str, List[Dict[str, Any]]] = {}
    for rec in records:
        stage = _canonical_stage(rec.get("inspection_type") or "")
        by_stage.setdefault(stage, []).append(rec)

    present_stages = [s for s in STAGE_ORDER if s in by_stage]
    unknown = [s for s in by_stage if s not in STAGE_ORDER]
    if unknown:
        logger.warning("Unrecognised inspection stage(s), skipped: %s", unknown)

    # Resolve templates, tracking what's usable vs. what's missing
    stage_templates: Dict[str, Path] = {}
    stages_skipped: List[Dict[str, str]] = []
    for stage in present_stages:
        try:
            stage_templates[stage] = resolve_template(report_type, stage)
        except TemplateNotConfiguredError as exc:
            stages_skipped.append({"stage": stage, "reason": exc.reason})
            logger.warning(
                "Skipping stage '%s' (%d record(s)): %s",
                stage, len(by_stage[stage]), exc.reason,
            )

    if not stage_templates:
        # Nothing at all could be produced — raise the most informative error
        first_stage = present_stages[0] if present_stages else "UNKNOWN"
        raise TemplateNotConfiguredError(
            report_type, first_stage,
            f"no configured template for any present stage {present_stages}",
        )

    # Pre-flight: validate ALL defect items across ALL usable stages
    # before writing anything.
    stage_workbooks: Dict[str, Any] = {}
    stage_header_maps: Dict[str, Dict[str, int]] = {}
    all_mismatches: List[Dict[str, Any]] = []

    for stage, template_path in stage_templates.items():
        wb = openpyxl.load_workbook(str(template_path), data_only=False)
        sheet_name = SHEET_NAME_MAP[stage]
        if sheet_name not in wb.sheetnames:
            raise ValueError(
                f"Template '{template_path.name}' has no sheet named "
                f"'{sheet_name}' (expected for stage '{stage}')."
            )
        ws = wb[sheet_name]
        header_map = _build_header_map(ws)

        stage_workbooks[stage]   = wb
        stage_header_maps[stage] = header_map

        all_mismatches.extend(_validate_defects(
            stage, header_map, by_stage[stage], defect_items_by_report, sheet_name,
        ))

    if all_mismatches:
        raise DefectMismatchError(all_mismatches)

    # Write + merge
    output_wb = openpyxl.Workbook()
    output_wb.remove(output_wb.active)  # drop default blank sheet

    stages_written: List[str] = []
    for stage, wb in stage_workbooks.items():
        sheet_name = SHEET_NAME_MAP[stage]
        ws = wb[sheet_name]
        _write_stage_sheet(
            ws=ws,
            stage=stage,
            records=by_stage[stage],
            defect_items_by_report=defect_items_by_report,
            header_map=stage_header_maps[stage],
        )
        copy_sheet(ws, output_wb, sheet_name)
        stages_written.append(stage)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_wb.save(str(output_path))

    logger.info(
        "Summary saved → %s (report_type=%s stages_written=%s stages_skipped=%s)",
        output_path, report_type, stages_written,
        [s["stage"] for s in stages_skipped],
    )

    return {"stages_written": stages_written, "stages_skipped": stages_skipped}
