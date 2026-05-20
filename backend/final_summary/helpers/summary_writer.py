"""
final_summary/helpers/summary_writer.py
----------------------------------------
Fills a pre-built Excel template with audit data.

Key behaviours
--------------
- Defect columns are placed by matching item names against the template's
  own header row (DEFECT_HEADER_ROW). The column index from the template
  is always authoritative — position is never assumed or offset-based.
- If ANY defect item in the data has no matching column in the template,
  a DefectMismatchError is raised BEFORE any file is written. The error
  carries a structured list of all mismatches so the caller can return a
  JSON error response to the user (same pattern as the Retry flow).
- Case-insensitive match is accepted as a valid match (no error).
- Only cell.value is written — no font / border / fill / alignment is
  touched so the template's pre-set row styles are fully preserved.
- Prefix columns: any key absent from the record dict is silently skipped
  (formula columns, future columns, etc.). Comment out entries in
  PREFIX_COLUMNS / INLINE_PREFIX_COLUMNS to skip intentionally.
- Values are normalised to native Python types:
    str   → factory, inspection_type, audit_result, report_no, item_name,
            style_no, po_no, country, acceptable_defect_qty, person
    date  → date_of_issue, po_wh, exf, po_edt, plan_edt, plan_wh
    time  → factory_in, factory_out, audit_start, audit_end
    float → defect_percentage
    int   → ship_qty, audit_qty, defect_qty, po_qty_pcs/pack/set, do_qty
"""

import logging
import re
from datetime import date, datetime, time
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import openpyxl
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)


# =============================================================================
# ── LAYOUT CONSTANTS  (edit here to adjust template positions) ───────────────
# =============================================================================

FINAL_DATA_START_ROW:  int = 12   # Final and Re-Final sheets
INLINE_DATA_START_ROW: int = 5    # INLINE sheet (and future SAMPLE)

FINAL_DEFECT_START_COL:  int = 26   # col Z
INLINE_DEFECT_START_COL: int = 21   # col U

DEFECT_HEADER_ROW: int = 2          # row in template that holds defect names


# =============================================================================
# ── PREFIX COLUMNS  (Final / Re-Final)
# Format: (column_index_1based, record_dict_key)
# Comment out formula-driven columns — they are skipped with no error.
# =============================================================================

PREFIX_COLUMNS: List[Tuple[int, str]] = [
    (1,  "factory"),
    (2,  "date_of_issue"),
    (3,  "inspection_type"),
    (4,  "factory_in"),
    (5,  "factory_out"),
    # (6,  "factory_total_hours"),   # formula column — skip
    (7,  "audit_start"),
    (8,  "audit_end"),
    # (9,  "audit_total_hours"),     # formula column — skip
    (10, "audit_result"),
    (11, "report_no"),
    (12, "item_name"),
    (13, "style_no"),
    (14, "po_no"),
    # (15, "country"),
    (16, "po_qty_pcs"),
    (17, "po_qty_pack"),
    (18, "po_qty_set"),
    (19, "po_wh"),
    (20, "ship_qty"),
    (21, "audit_qty"),
    (22, "acceptable_defect_qty"),
    (23, "defect_qty"),
    (24, "defect_percentage"),
    (25, "person"),
]

# =============================================================================
# ── INLINE PREFIX COLUMNS  (skeleton — update when INLINE template confirmed)
# Defects start at col U (21) so prefix fits in cols A–T (1–20).
# =============================================================================

INLINE_PREFIX_COLUMNS: List[Tuple[int, str]] = [
    # TODO: confirm exact INLINE prefix column order with template owner
    (1,  "factory"),
    (2,  "date_of_issue"),
    (3,  "inspection_type"),
    (4,  "factory_in"),
    (5,  "factory_out"),
    # (6,  "factory_total_hours"),   # formula column — skip
    (7,  "audit_start"),
    (8,  "audit_end"),
    # (9,  "audit_total_hours"),     # formula column — skip
    (10, "audit_result"),
    (11, "report_no"),
    (12, "item_name"),
    (13, "style_no"),
    (14, "po_no"),
    # (15, "country"),
    (16, "ship_qty"),
    (17, "audit_qty"),
    (18, "defect_qty"),
    (19, "defect_percentage"),
    (20, "person"),
]

# =============================================================================
# ── SHEET NAME MAP ───────────────────────────────────────────────────────────
# =============================================================================

SHEET_NAME_MAP: Dict[str, str] = {
    "FINAL":    "Final",
    "RE-FINAL": "Re-Final",
    "INLINE":   "INLINE",
    # "SAMPLE": "Sample",   # add when template sheet is ready
}

INLINE_CANONICAL_TYPES: set = {"INLINE"}

# =============================================================================
# ── FIELD TYPE DECLARATIONS ───────────────────────────────────────────────────
# =============================================================================

_STR_FIELDS: set = {
    "factory", "inspection_type", "audit_result", "report_no",
    "item_name", "style_no", "po_no", "country", "acceptable_defect_qty",
    "person", "inspector",
}
_DATE_FIELDS: set = {
    "date_of_issue", "po_wh", "exf", "po_edt", "plan_edt", "plan_wh",
}
_TIME_FIELDS: set = {
    "factory_in", "factory_out", "audit_start", "audit_end",
}
_FLOAT_FIELDS: set = {"defect_percentage"}
_INT_FIELDS: set = {
    "ship_qty", "audit_qty", "defect_qty", "do_qty",
    "po_qty_pcs", "po_qty_pack", "po_qty_set",
}


# =============================================================================
# ── CUSTOM EXCEPTION ─────────────────────────────────────────────────────────
# =============================================================================

class DefectMismatchError(Exception):
    """
    Raised when one or more defect items in the data have no matching column
    in the template header row.  Carries a structured list of mismatches so
    the caller can return a JSON error response.

    Attributes
    ----------
    mismatches : list of dicts, each with keys:
        file_name   : source file name
        report_id   : AuditReport pk
        sheet       : template sheet name where the match was attempted
        item        : defect item name that was not found
        category    : defect category
    """

    def __init__(self, mismatches: List[Dict[str, Any]]) -> None:
        self.mismatches = mismatches
        count = len(mismatches)
        super().__init__(
            f"{count} defect item(s) have no matching column in the template."
        )


# =============================================================================
# ── VALUE NORMALISATION ───────────────────────────────────────────────────────
# =============================================================================

def _to_str(value: Any) -> Optional[str]:
    if value is None or value == "":
        return None
    s = str(value).strip()
    return s if s and s != "-" else None


def _to_date(value: Any) -> Optional[date]:
    if value is None or value == "":
        return None
    if isinstance(value, date):
        return value
    s = str(value).strip()
    for fmt in (
        "%Y-%m-%d", "%m/%d/%Y", "%d/%m/%Y",
        "%Y/%m/%d", "%m-%d-%Y", "%d-%m-%Y",
        "%d-%b-%y", "%d-%b-%Y",
    ):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    logger.debug("_to_date: unrecognised format '%s'", s)
    return None


def _to_time(value: Any) -> Optional[time]:
    if value is None or value == "":
        return None
    if isinstance(value, time):
        return value
    s = str(value).strip()
    m = re.match(r"(\d{1,2})[:.](\d{2})(?:\s*([AaPp][Mm]))?", s)
    if not m:
        return None
    hour, minute, ampm = int(m.group(1)), int(m.group(2)), m.group(3)
    if ampm:
        if ampm.upper() == "PM" and hour != 12:
            hour += 12
        elif ampm.upper() == "AM" and hour == 12:
            hour = 0
    try:
        return time(hour, minute)
    except ValueError:
        return None


def _to_float(value: Any) -> Optional[float]:
    if value is None or value == "":
        return None
    try:
        return float(str(value).replace("%", "").strip())
    except (ValueError, TypeError):
        return None


def _to_int(value: Any) -> Optional[int]:
    if value is None or value == "":
        return None
    try:
        return int(float(str(value).strip()))
    except (ValueError, TypeError):
        return None


def _normalise(key: str, value: Any) -> Any:
    if key in _STR_FIELDS:
        return _to_str(value)
    if key in _DATE_FIELDS:
        return _to_date(value)
    if key in _TIME_FIELDS:
        return _to_time(value)
    if key in _FLOAT_FIELDS:
        return _to_float(value)
    if key in _INT_FIELDS:
        return _to_int(value)
    if value == "" or value == "-":
        return None
    return value


# =============================================================================
# ── DEFECT HEADER MAP ─────────────────────────────────────────────────────────
# =============================================================================

def _build_defect_header_map(ws, defect_start_col: int) -> Dict[str, int]:
    """
    Read DEFECT_HEADER_ROW from the template and return
    {item_name_stripped: column_index_1based}.

    Scans left-to-right from defect_start_col. First occurrence wins.
    Blank cells are skipped.
    """
    header_map: Dict[str, int] = {}
    for col in range(defect_start_col, ws.max_column + 1):
        raw = ws.cell(row=DEFECT_HEADER_ROW, column=col).value
        if raw is None:
            continue
        name = str(raw).strip()
        if name and name not in header_map:
            header_map[name] = col

    logger.debug(
        "_build_defect_header_map: sheet='%s' found %d headers from col %s",
        ws.title, len(header_map), get_column_letter(defect_start_col),
    )
    return header_map


def _resolve_defect_col(
    item_name: str,
    header_map: Dict[str, int],
) -> Optional[int]:
    """
    1. Exact (case-sensitive) match.
    2. Case-insensitive match.
    3. None → item not in template (caller must raise error).
    """
    col = header_map.get(item_name)
    if col is not None:
        return col
    lower = item_name.lower()
    for name, c in header_map.items():
        if name.lower() == lower:
            return c
    return None


# =============================================================================
# ── CANONICAL TYPE HELPERS ────────────────────────────────────────────────────
# =============================================================================

def _canonical_type(itype: str) -> str:
    u = (itype or "").upper().strip()
    if re.search(r"PRE[\s\-]?FINAL", u):                          return "PRE-FINAL"
    if re.search(r"(?<![A-Z])RE[\s\-]?FINAL|REFINAL", u):        return "RE-FINAL"
    if re.search(r"\bFINAL\b|SHIPMENT\s*AUDIT|PRE[\s\-]?SHIP", u): return "FINAL"
    if re.search(r"IN[\s\-]?LINE", u):                            return "INLINE"
    if re.search(r"\bCMF\b|COUNTER\s*MASTER", u):                 return "CMF"
    if re.search(r"\bSAMPLE\b|PRE[\s\-]?PROD|\bPP\b", u):        return "SAMPLE"
    return "UNKNOWN"


def _is_inline(canonical: str) -> bool:
    return canonical in INLINE_CANONICAL_TYPES


def _defect_start_col(canonical: str) -> int:
    return INLINE_DEFECT_START_COL if _is_inline(canonical) else FINAL_DEFECT_START_COL


def _data_start_row(canonical: str) -> int:
    return INLINE_DATA_START_ROW if _is_inline(canonical) else FINAL_DATA_START_ROW


def _prefix_columns(canonical: str) -> List[Tuple[int, str]]:
    return INLINE_PREFIX_COLUMNS if _is_inline(canonical) else PREFIX_COLUMNS


# =============================================================================
# ── PRE-FLIGHT DEFECT VALIDATION ──────────────────────────────────────────────
# =============================================================================

def _validate_defects(
    wb,
    by_canonical: Dict[str, List[Dict[str, Any]]],
    defect_items_by_report: Dict[int, List[Dict[str, Any]]],
) -> List[Dict[str, Any]]:
    """
    Check every defect item for every record against the template header map.
    Returns a flat list of mismatch dicts (empty = all clear).

    Runs BEFORE any cell is written so the file is never partially filled.
    """
    mismatches: List[Dict[str, Any]] = []

    for canonical, records in by_canonical.items():
        sheet_name = SHEET_NAME_MAP.get(canonical)
        if not sheet_name or sheet_name not in wb.sheetnames:
            # Sheet-level issues are handled separately; skip here
            continue

        ws           = wb[sheet_name]
        defect_start = _defect_start_col(canonical)
        header_map   = _build_defect_header_map(ws, defect_start)

        for rec in records:
            rid       = rec.get("report_id")
            file_name = rec.get("file_name", "")

            for d in defect_items_by_report.get(rid, []):
                item_name = (d.get("item") or "").strip()
                if not item_name:
                    continue
                major = _to_int(d.get("major_count", 0))
                if not major:
                    continue  # zero-count rows — no column needed

                if _resolve_defect_col(item_name, header_map) is None:
                    mismatches.append({
                        "file_name": file_name,
                        "report_id": rid,
                        "sheet":     sheet_name,
                        "item":      item_name,
                        "category":  (d.get("category") or "").strip(),
                    })

    return mismatches


# =============================================================================
# ── SHEET WRITER ──────────────────────────────────────────────────────────────
# =============================================================================

def _write_sheet(
    ws,
    records: List[Dict[str, Any]],
    defect_items_by_report: Dict[int, List[Dict[str, Any]]],
    canonical: str,
    header_map: Dict[str, int],
) -> int:
    """
    Fill one template sheet. header_map is pre-built and validated.
    Only cell.value is set — no styling is touched.
    Returns number of rows written.
    """
    start_row   = _data_start_row(canonical)
    prefix_cols = _prefix_columns(canonical)
    current_row = start_row

    for rec in records:
        rid = rec.get("report_id")

        # ── Prefix columns ────────────────────────────────────────────────────
        for col_idx, key in prefix_cols:
            raw = rec.get(key)
            if raw is None:
                continue  # absent key = formula/skip column — leave untouched
            value = _normalise(key, raw)
            if value is not None:
                ws.cell(row=current_row, column=col_idx).value = value

        # ── Defect columns ────────────────────────────────────────────────────
        for d in defect_items_by_report.get(rid, []):
            item_name = (d.get("item") or "").strip()
            if not item_name:
                continue
            major = _to_int(d.get("major_count", 0))
            if not major:
                continue  # zero count — leave cell blank (formula still works)

            col_idx = _resolve_defect_col(item_name, header_map)
            # col_idx is guaranteed non-None here (validated in pre-flight)
            ws.cell(row=current_row, column=col_idx).value = major

        current_row += 1

    rows_written = current_row - start_row
    logger.info(
        "Sheet '%s': wrote %d row(s) | start_row=%d | defect_col_start=%s",
        ws.title, rows_written, start_row,
        get_column_letter(_defect_start_col(canonical)),
    )
    return rows_written


# =============================================================================
# ── PUBLIC API ────────────────────────────────────────────────────────────────
# =============================================================================

def write_summary(
    records: List[Dict[str, Any]],
    defect_items_by_report: Dict[int, List[Dict[str, Any]]],
    output_path: Path,
    template_path: Path,
    format_type: str = "SPI",
) -> None:
    """
    Fill the pre-built Excel template and save to output_path.

    Raises
    ------
    FileNotFoundError   if template_path does not exist
    DefectMismatchError if any defect item has no matching template column;
                        inspect .mismatches for the full structured list
    ValueError          if no records could be written to any sheet

    Parameters
    ----------
    records                : list of record dicts (one per AuditReport row)
    defect_items_by_report : {report_id: [{category, item, major_count, …}]}
    output_path            : destination .xlsx (parent dir created if needed)
    template_path          : absolute Path to the pre-built template .xlsx
    format_type            : "SPI" | "REGULAR" | "SWEATER"
    """
    if not template_path.exists():
        raise FileNotFoundError(f"Template not found: {template_path}")

    wb = openpyxl.load_workbook(str(template_path), data_only=False)

    # Group by canonical type
    by_canonical: Dict[str, List[Dict[str, Any]]] = {}
    for rec in records:
        key = _canonical_type(rec.get("inspection_type") or "")
        by_canonical.setdefault(key, []).append(rec)

    # ── Pre-flight: validate all defect items against template headers ─────────
    # Must happen BEFORE any cell is written so the file is never half-filled.
    mismatches = _validate_defects(wb, by_canonical, defect_items_by_report)
    if mismatches:
        raise DefectMismatchError(mismatches)

    # ── Write sheets ──────────────────────────────────────────────────────────
    process_order = ["FINAL", "RE-FINAL", "INLINE"] + [
        k for k in by_canonical if k not in ("FINAL", "RE-FINAL", "INLINE")
    ]

    sheets_written = 0
    skipped_types: List[str] = []

    for canonical in process_order:
        type_records = by_canonical.get(canonical)
        if not type_records:
            continue

        sheet_name = SHEET_NAME_MAP.get(canonical)
        if not sheet_name:
            skipped_types.append(canonical)
            logger.warning(
                "No template sheet mapped for type '%s' (%d record(s) skipped). "
                "Add to SHEET_NAME_MAP and create the sheet in the template.",
                canonical, len(type_records),
            )
            continue

        if sheet_name not in wb.sheetnames:
            skipped_types.append(canonical)
            logger.warning(
                "Sheet '%s' not found in template '%s' — "
                "skipping %d record(s) of type '%s'.",
                sheet_name, template_path.name, len(type_records), canonical,
            )
            continue

        ws           = wb[sheet_name]
        defect_start = _defect_start_col(canonical)
        header_map   = _build_defect_header_map(ws, defect_start)

        _write_sheet(
            ws=ws,
            records=type_records,
            defect_items_by_report=defect_items_by_report,
            canonical=canonical,
            header_map=header_map,
        )
        sheets_written += 1

    if not sheets_written:
        raise ValueError(
            "No records written — no canonical inspection types matched "
            "a sheet in the template."
        )

    if skipped_types:
        logger.warning(
            "Inspection types skipped (no matching template sheet): %s",
            skipped_types,
        )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output_path))
    logger.info(
        "Summary saved → %s  (template=%s  format=%s  sheets=%d)",
        output_path, template_path.name, format_type, sheets_written,
    )


# =============================================================================
# ── LEGACY HELPER ─────────────────────────────────────────────────────────────
# =============================================================================

def build_defect_column_plan(
    defect_rows_by_report: Dict[int, List[Dict[str, Any]]],
) -> List[Tuple[str, str]]:
    seen: Dict[Tuple[str, str], None] = {}
    for rows in defect_rows_by_report.values():
        for d in rows:
            if isinstance(d, dict):
                cat  = (d.get("category") or "").strip()
                item = (d.get("item") or "").strip()
                if cat and item:
                    seen[(cat, item)] = None
    return sorted(seen.keys(), key=lambda x: (x[0], x[1]))