import openpyxl
import logging
from pathlib import Path
logger = logging.getLogger(__name__)

def _convert_xls_to_xlsx(xls_path: Path) -> Path:
    """
    Convert .xls to .xlsx in the same temp folder.
    Returns the path to the converted .xlsx file.
    """
    import xlrd
    from xlrd import xldate_as_datetime

    xlsx_path = xls_path.with_suffix(".xlsx")
    xls_wb    = xlrd.open_workbook(str(xls_path), formatting_info=False)
    xlsx_wb   = openpyxl.Workbook()
    xlsx_wb.remove(xlsx_wb.active)  # remove default blank sheet

    for sheet_idx in range(xls_wb.nsheets):
        xls_ws  = xls_wb.sheet_by_index(sheet_idx)
        xlsx_ws = xlsx_wb.create_sheet(title=xls_ws.name)
        for row_idx in range(xls_ws.nrows):
            for col_idx in range(xls_ws.ncols):
                cell = xls_ws.cell(row_idx, col_idx)
                if cell.ctype == xlrd.XL_CELL_DATE:
                    value = xldate_as_datetime(cell.value, xls_wb.datemode)
                else:
                    value = cell.value
                xlsx_ws.cell(row=row_idx + 1, column=col_idx + 1).value = value

    xlsx_wb.save(str(xlsx_path))
    xls_wb.release_resources()
    logger.info("Converted %s → %s", xls_path.name, xlsx_path.name)
    return xlsx_path
